#!/usr/bin/env python3
"""
Import V2 rows from a flat Excel sheet into data/si-schedule-db.json (scope v2).

Expected columns: Date, time, topic, intro, Atten, (division), loc
Handles group matrices (2- and 3-column) and parallel sessions within the same time block.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError as e:  # pragma: no cover
    print("Need openpyxl: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1) from e

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data" / "si-schedule-db.json"
DEFAULT_XLSX = Path("/Users/jo/Desktop/5月/V2.xlsx")

DAY_ID_BY_CAL: dict[int, str] = {3: "d1", 4: "d2", 5: "d3", 6: "d4", 7: "d5"}

CAMPUS_CODES = frozenset(
    {
        "BISZ",
        "BIGZ",
        "BIHZ",
        "BIPH",
        "BINJ",
        "BBSZ",
        "BIBCD",
        "BIBWH",
        "BBGM",
        "BKNS",
        "BKFT",
        "BBGZ",
        "BBBJ",
    }
)

DIVISION_LABELS = frozenset({"School", "ECE", "PS", "MS", "HS", "MS/HS"})


def cell_str(v: object, *, flatten_newlines: bool = False) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, dt.datetime):
        return v.strftime("%H:%M")
    s = str(v).strip()
    if flatten_newlines:
        s = re.sub(r"\s*\n\s*", ", ", s)
    return s


def norm_time(v: object) -> str:
    s = cell_str(v)
    if not s:
        return ""
    return re.sub(r"\s+", " ", s)


def parse_day_id(date_cell: object) -> str | None:
    s = cell_str(date_cell)
    if not s:
        return None
    compact = re.sub(r"\s+", "", s)
    m = re.search(r"8月(\d{1,2})日", compact)
    if not m:
        m = re.search(r"Aug(?:ust)?\s*(\d{1,2})", s, re.I)
    if not m:
        return None
    return DAY_ID_BY_CAL.get(int(m.group(1)))


def empty_row() -> dict[str, str]:
    return {
        "timeZh": "",
        "timeEn": "",
        "topicZh": "",
        "topicEn": "",
        "introZh": "",
        "introEn": "",
        "locZh": "",
        "locEn": "",
        "detailZh": "",
        "detailEn": "",
        "groupTsvZh": "",
        "groupTsvEn": "",
    }


def is_group_label(atten: str) -> bool:
    return bool(re.match(r"^Group\s+\d+$", atten, re.I))


def is_campus(atten: str) -> bool:
    return atten in CAMPUS_CODES


def is_division(col6: str) -> bool:
    return col6 in DIVISION_LABELS


def matrix_to_tsv(kind: str, rows: list[list[str]]) -> str:
    if not rows:
        return ""
    if kind == "3col":
        header = "Campus\tGroup\tLocation"
    else:
        header = "Attendees\tLocation"
    lines = [header]
    for r in rows:
        lines.append("\t".join(r))
    return "\n".join(lines)


class MatrixAcc:
    def __init__(self) -> None:
        self.kind: str | None = None  # "2col" | "3col"
        self.rows: list[list[str]] = []
        self.campus: str = ""

    def clear(self) -> None:
        self.kind = None
        self.rows = []
        self.campus = ""

    def add_2col(self, atten: str, loc: str) -> None:
        if self.kind == "3col" and self.rows:
            return
        self.kind = "2col"
        if atten or loc:
            self.rows.append([atten, loc])

    def add_3col(self, campus: str, group: str, loc: str) -> None:
        if self.kind == "2col" and self.rows:
            return
        self.kind = "3col"
        if campus:
            self.campus = campus
        c = campus or self.campus
        if c and group and loc:
            self.rows.append([c, group, loc])
        elif c and group and not loc:
            self.rows.append([c, group, ""])
        elif not c and group and loc and self.campus:
            self.rows.append([self.campus, group, loc])

    def flush_to(self, entry: dict[str, str] | None) -> None:
        if entry and self.rows and self.kind:
            entry["groupTsvZh"] = matrix_to_tsv(self.kind, self.rows)
            entry["groupTsvEn"] = entry["groupTsvZh"]
            entry["locZh"] = ""
            entry["locEn"] = ""
        self.clear()


def uses_group_matrix(atten: str, loc: str) -> bool:
    if not atten or atten.upper() == "ALL" or not loc:
        return False
    return is_group_label(atten) or is_campus(atten)


def atten_to_intro(atten: str) -> str:
    if atten.upper() == "ALL":
        return "ALL"
    return ""


def should_start_parallel(topic: str, intro: str, atten: str, col6: str, loc: str, has_time: bool) -> bool:
    if has_time:
        return bool(topic)
    if not topic:
        return False
    if intro:
        return True
    if atten and not is_group_label(atten) and not is_campus(atten):
        return True
    if atten and loc:
        return True
    if atten and is_group_label(atten):
        return True
    if col6 and not is_division(col6):
        return True
    return False


def is_matrix_row(topic: str, atten: str, col6: str, loc: str) -> bool:
    if topic:
        return False
    if not atten and not col6 and not loc:
        return False
    if is_campus(atten) or (not atten and col6 and is_division(col6)):
        return True
    if is_group_label(atten):
        return True
    if atten and loc and not is_campus(atten):
        return True
    if atten and not loc and is_group_label(atten):
        return True
    return False


def read_v2_rows(path: Path) -> dict[str, list[dict[str, str]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = "Sheet1" if "Sheet1" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]

    by_day: dict[str, list[dict[str, str]]] = {f"d{i}": [] for i in range(1, 6)}
    day_id: str | None = None
    current_time = ""
    last_entry: dict[str, str] | None = None
    matrix = MatrixAcc()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        cells = list(row) + [None] * max(0, 7 - len(row))
        date_c, time_c, topic_c, intro_c, atten_c, col6_c, loc_c = cells[:7]

        parsed_day = parse_day_id(date_c)
        if parsed_day:
            matrix.flush_to(last_entry)
            day_id = parsed_day
            current_time = ""
            last_entry = None

        if not day_id:
            continue

        topic = cell_str(topic_c)
        intro = cell_str(intro_c)
        atten = cell_str(atten_c, flatten_newlines=True)
        col6 = cell_str(col6_c, flatten_newlines=True)
        loc = cell_str(loc_c, flatten_newlines=True)
        time_s = norm_time(time_c)
        has_time = bool(time_s)

        if has_time:
            matrix.flush_to(last_entry)
            current_time = time_s

        if has_time and topic:
            entry = empty_row()
            entry["timeZh"] = current_time
            entry["timeEn"] = current_time
            entry["topicZh"] = topic
            entry["topicEn"] = topic
            entry["introZh"] = intro or atten_to_intro(atten)
            entry["introEn"] = entry["introZh"]
            if loc and not uses_group_matrix(atten, loc):
                entry["locZh"] = loc
                entry["locEn"] = loc
            by_day[day_id].append(entry)
            last_entry = entry
            if uses_group_matrix(atten, loc):
                if is_campus(atten) and col6 and is_division(col6):
                    matrix.add_3col(atten, col6, loc)
                else:
                    matrix.add_2col(atten, loc)
            continue

        if has_time and not topic:
            continue

        if topic and should_start_parallel(topic, intro, atten, col6, loc, has_time=False):
            matrix.flush_to(last_entry)
            entry = empty_row()
            entry["timeZh"] = current_time
            entry["timeEn"] = current_time
            entry["topicZh"] = topic
            entry["topicEn"] = topic
            entry["introZh"] = intro
            entry["introEn"] = intro
            if loc and not uses_group_matrix(atten, loc):
                entry["locZh"] = loc
                entry["locEn"] = loc
            elif atten and not is_group_label(atten) and not is_campus(atten):
                entry["introZh"] = intro or atten
                entry["introEn"] = entry["introZh"]
            by_day[day_id].append(entry)
            last_entry = entry
            if uses_group_matrix(atten, loc):
                if is_campus(atten) and col6 and is_division(col6):
                    matrix.add_3col(atten, col6, loc)
                else:
                    matrix.add_2col(atten, loc)
            continue

        if topic and last_entry and not atten and not loc and not col6 and not intro:
            prev = last_entry["topicZh"]
            last_entry["topicZh"] = (prev + "\n" + topic).strip() if prev else topic
            last_entry["topicEn"] = last_entry["topicZh"]
            if intro:
                prev_i = last_entry["introZh"]
                last_entry["introZh"] = (prev_i + "\n" + intro).strip() if prev_i else intro
                last_entry["introEn"] = last_entry["introZh"]
            continue

        if is_matrix_row(topic, atten, col6, loc):
            if is_campus(atten):
                if col6 and is_division(col6):
                    matrix.add_3col(atten, col6, loc)
                else:
                    matrix.add_2col(atten, loc)
            elif not atten and col6 and is_division(col6):
                matrix.add_3col("", col6, loc)
            elif is_group_label(atten):
                matrix.add_2col(atten, loc)
            elif atten and loc:
                matrix.add_2col(atten, loc)
            elif atten and not loc and last_entry:
                note = atten
                prev = last_entry["introZh"]
                last_entry["introZh"] = (prev + "\n" + note).strip() if prev else note
                last_entry["introEn"] = last_entry["introZh"]
            continue

        if atten and not topic and not loc and last_entry:
            prev = last_entry["introZh"]
            last_entry["introZh"] = (prev + "\n" + atten).strip() if prev else atten
            last_entry["introEn"] = last_entry["introZh"]

    matrix.flush_to(last_entry)
    wb.close()
    return by_day


def main() -> None:
    ap = argparse.ArgumentParser(description="Import V2 schedule from Excel into si-schedule-db.json")
    ap.add_argument("xlsx", type=Path, nargs="?", default=DEFAULT_XLSX)
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    xlsx = args.xlsx.expanduser().resolve()
    out_path = args.output.expanduser().resolve()
    if not xlsx.is_file():
        raise SystemExit(f"Not a file: {xlsx}")

    by_day = read_v2_rows(xlsx)

    if out_path.is_file():
        doc = json.loads(out_path.read_text(encoding="utf-8"))
    else:
        doc = {"v": 1, "entries": {}}
    entries = doc.get("entries")
    if not isinstance(entries, dict):
        entries = {}

    for k in list(entries.keys()):
        if k.startswith("v2:"):
            del entries[k]
    for day_id, arr in by_day.items():
        entries["v2:" + day_id] = arr

    doc["entries"] = entries
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    total = sum(len(v) for v in by_day.values())
    counts = ", ".join(f"{k}={len(v)}" for k, v in sorted(by_day.items()) if v)
    print(f"Wrote {out_path}")
    print(f"V2 rows: {total} ({counts})")


if __name__ == "__main__":
    main()
