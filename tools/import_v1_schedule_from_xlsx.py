#!/usr/bin/env python3
"""
Import V1 rows from a flat Excel «Schedule» sheet into data/si-schedule-db.json.

Expected columns (header row): Date, time, topic, loc, intro
- Date: e.g. «8月 3 日» … «8月 7 日» → day_id d1 … d5 (Aug 3–7, 2026)
- Other columns map to schedule-db row fields (zh/en duplicated when only one language).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError as e:  # pragma: no cover
    print("Need openpyxl: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1) from e

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data" / "si-schedule-db.json"

# Calendar day in August 2026 → d1..d5 (institute week)
DAY_ID_BY_CAL: dict[int, str] = {3: "d1", 4: "d2", 5: "d3", 6: "d4", 7: "d5"}


def cell_str(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s


def parse_day_id(date_cell: object) -> str | None:
    s = cell_str(date_cell)
    if not s:
        return None
    compact = re.sub(r"\s+", "", s)
    m = re.search(r"8月(\d{1,2})日", compact)
    if not m:
        m = re.search(r"Aug(?:ust)?\s*(\d{1,2})", s, re.I)
    if not m:
        return None
    cal = int(m.group(1))
    return DAY_ID_BY_CAL.get(cal)


def norm_header(h: object) -> str:
    return re.sub(r"\s+", "", str(h or "").strip().lower())


def row_dict(
    time_s: str,
    topic_s: str,
    loc_s: str,
    intro_s: str,
) -> dict[str, str]:
    """One DB row; mirror emptyRow() in schedule-db-page.js."""
    t = time_s.strip()
    topic = topic_s.strip()
    loc = loc_s.strip()
    intro = intro_s.strip()
    return {
        "timeZh": t,
        "timeEn": t,
        "topicZh": topic,
        "topicEn": topic,
        "introZh": intro,
        "introEn": intro,
        "locZh": loc,
        "locEn": loc,
        "detailZh": "",
        "detailEn": "",
        "groupTsvZh": "",
        "groupTsvEn": "",
    }


def read_schedule_rows(path: Path) -> list[tuple[str | None, dict[str, str]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Schedule" not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"No sheet «Schedule»; found: {wb.sheetnames}")
    ws = wb["Schedule"]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        wb.close()
        return []

    idx: dict[str, int] = {}
    for i, h in enumerate(header):
        key = norm_header(h)
        if key in ("date", "time", "topic", "loc", "intro"):
            idx[key] = i
    required = ("date", "time", "topic")
    missing = [k for k in required if k not in idx]
    if missing:
        wb.close()
        raise SystemExit(f"Missing columns {missing}; header={header!r}")

    out: list[tuple[str | None, dict[str, str]]] = []
    for row in rows_iter:
        if not row:
            continue
        def get(name: str) -> str:
            j = idx.get(name)
            if j is None or j >= len(row):
                return ""
            return cell_str(row[j])

        day_id = parse_day_id(get("date"))
        time_s = get("time")
        topic_s = get("topic")
        if not day_id and not time_s and not topic_s:
            continue
        if not day_id:
            print(f"Skip row (unparsed date): {row!r}", file=sys.stderr)
            continue
        if not topic_s and not time_s:
            continue
        out.append(
            (
                day_id,
                row_dict(
                    time_s,
                    topic_s,
                    get("loc"),
                    get("intro"),
                ),
            )
        )
    wb.close()
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="Import V1 schedule from Excel into si-schedule-db.json")
    ap.add_argument(
        "xlsx",
        type=Path,
        nargs="?",
        default=Path("/Users/jo/Desktop/5月/夏训/05 文件/V1-schedule.xlsx"),
        help="Path to V1-schedule.xlsx",
    )
    ap.add_argument(
        "-o",
        "--output",
        type=Path,
        default=DEFAULT_OUT,
        help=f"Output JSON (default: {DEFAULT_OUT})",
    )
    args = ap.parse_args()
    xlsx: Path = args.xlsx.expanduser().resolve()
    out_path: Path = args.output.expanduser().resolve()
    if not xlsx.is_file():
        raise SystemExit(f"Not a file: {xlsx}")

    pairs = read_schedule_rows(xlsx)
    by_day: dict[str, list[dict[str, str]]] = {f"d{i}": [] for i in range(1, 6)}
    for day_id, rec in pairs:
        by_day.setdefault(day_id, []).append(rec)

    if out_path.is_file():
        doc = json.loads(out_path.read_text(encoding="utf-8"))
    else:
        doc = {"v": 1, "entries": {}}
    if not isinstance(doc, dict) or "entries" not in doc:
        doc = {"v": int(doc.get("v", 1)) if isinstance(doc, dict) else 1, "entries": {}}
    entries = doc.get("entries")
    if not isinstance(entries, dict):
        entries = {}

    for k in list(entries.keys()):
        if k.startswith("v1:"):
            del entries[k]
    for day_id, arr in by_day.items():
        if arr:
            entries["v1:" + day_id] = arr

    doc["entries"] = entries
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    counts = ", ".join(f"{k}={len(v)}" for k, v in sorted(by_day.items()) if v)
    print(f"Wrote {out_path}")
    print(f"V1 rows: {len(pairs)} ({counts})")


if __name__ == "__main__":
    main()
